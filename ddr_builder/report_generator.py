from __future__ import annotations

from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as RLImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from ddr_builder.schema import DDRIntermediate


def _styles():
    base = getSampleStyleSheet()
    normal = ParagraphStyle(
        name="DDRBody",
        parent=base["Normal"],
        fontSize=10,
        leading=14,
        alignment=TA_JUSTIFY,
        spaceAfter=6,
    )
    heading = ParagraphStyle(
        name="DDRHeading",
        parent=base["Heading1"],
        fontSize=14,
        leading=18,
        spaceBefore=12,
        spaceAfter=8,
        textColor=colors.HexColor("#1a365d"),
    )
    sub = ParagraphStyle(
        name="DDRSub",
        parent=base["Heading2"],
        fontSize=11,
        leading=14,
        spaceBefore=8,
        spaceAfter=4,
        textColor=colors.HexColor("#2c5282"),
    )
    bullet = ParagraphStyle(
        name="DDRBullet",
        parent=normal,
        leftIndent=18,
        bulletIndent=6,
    )
    return normal, heading, sub, bullet


def _escape(text: str) -> str:
    return (
        (text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _p(flow, text: str, style):
    if not (text or "").strip():
        text = "Not Available"
    flow.append(Paragraph(_escape(text), style))


def _bullets(flow, items: list[str], bullet_style, normal_style):
    if not items:
        _p(flow, "Not Available", normal_style)
        return
    for it in items:
        flow.append(Paragraph(f"• {_escape(it)}", bullet_style))


def build_ddr_pdf(ddr: DDRIntermediate, out_path: str | Path, max_image_width: float = 5.5 * inch, base_dir: str | Path | None = None):
    """Render DDRIntermediate to a client-style PDF."""
    out_path = Path(out_path).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    _base = Path(base_dir).resolve() if base_dir else Path.cwd()

    normal, heading, sub, bullet = _styles()
    story: list = []

    story.append(Paragraph(_escape("DETAILED DIAGNOSTIC REPORT (DDR)"), heading))
    story.append(Spacer(1, 0.15 * inch))

    story.append(Paragraph(_escape("1. Property Issue Summary"), sub))
    _p(story, ddr.property_summary, normal)
    story.append(Spacer(1, 0.1 * inch))

    story.append(Paragraph(_escape("2. Area-wise Observations"), sub))
    for i, area in enumerate(ddr.areas):
        title = area.area_name or "Not Available"
        story.append(Paragraph(_escape(f"Area: {title}"), normal))
        _p(story, f"<b>Observation:</b> {area.observation}", normal)
        _p(story, f"<b>Thermal Finding:</b> {area.thermal_finding}", normal)
        _p(story, f"<b>Probable Root Cause:</b> {area.root_cause}", normal)
        sev = area.severity or "Not Available"
        reason = area.severity_reason or "Not Available"
        _p(story, f"<b>Severity:</b> {sev} — <b>Reason:</b> {reason}", normal)
        _p(story, f"<b>Recommended Action:</b> {area.recommended_action}", normal)
        if (area.conflict_note or "").strip():
            _p(story, f"<b>Conflict:</b> {area.conflict_note}", normal)

        story.append(Paragraph(_escape("<b>Image:</b>"), normal))
        imgs_added = False
        for pth in area.image_paths:
            path = Path(pth)
            if not path.is_absolute():
                path = (_base / path).resolve()
            if path.is_file():
                try:
                    img = RLImage(str(path))
                    img.drawWidth = max_image_width
                    if getattr(img, "imageWidth", None) and img.imageWidth > 0:
                        img.drawHeight = img.imageHeight * max_image_width / img.imageWidth
                    else:
                        img.drawHeight = 3 * inch
                    story.append(img)
                    story.append(Spacer(1, 0.08 * inch))
                    imgs_added = True
                except Exception:
                    continue
        if not imgs_added:
            label = "Image Not Available"
            if area.image_paths:
                label += " (paths present but files not found or unsupported)"
            _p(story, label, normal)

        story.append(Spacer(1, 0.15 * inch))
        if i < len(ddr.areas) - 1:
            story.append(Spacer(1, 0.05 * inch))

    story.append(PageBreak())
    story.append(Paragraph(_escape("3. Probable Root Cause Summary"), sub))
    _bullets(story, ddr.root_causes_summary, bullet, normal)

    story.append(Paragraph(_escape("4. Severity Assessment Summary"), sub))
    story.append(Paragraph(_escape("High Severity Issues:"), normal))
    _bullets(story, ddr.high_severity, bullet, normal)
    story.append(Paragraph(_escape("Medium Severity Issues:"), normal))
    _bullets(story, ddr.medium_severity, bullet, normal)
    story.append(Paragraph(_escape("Low Severity Issues:"), normal))
    _bullets(story, ddr.low_severity, bullet, normal)

    story.append(Paragraph(_escape("5. Recommended Actions Summary"), sub))
    story.append(Paragraph(_escape("Immediate Actions:"), normal))
    _bullets(story, ddr.immediate_actions, bullet, normal)
    story.append(Paragraph(_escape("Preventive Actions:"), normal))
    _bullets(story, ddr.preventive_actions, bullet, normal)
    story.append(Paragraph(_escape("Further Investigation Required:"), normal))
    _bullets(story, ddr.further_investigation, bullet, normal)

    story.append(Paragraph(_escape("6. Additional Notes"), sub))
    _bullets(story, ddr.additional_notes, bullet, normal)
    if ddr.conflicts:
        story.append(Paragraph(_escape("Document-level conflicts:"), normal))
        _bullets(story, ddr.conflicts, bullet, normal)

    story.append(Paragraph(_escape("7. Missing or Unclear Information"), sub))
    _bullets(story, ddr.missing_information, bullet, normal)

    doc = SimpleDocTemplate(
        str(out_path),
        pagesize=letter,
        rightMargin=56,
        leftMargin=56,
        topMargin=56,
        bottomMargin=56,
    )
    doc.build(story)


def build_ddr_docx(
    ddr: DDRIntermediate,
    out_path: str | Path,
    max_image_width_inches: float = 5.5,
    base_dir: str | Path | None = None,
) -> None:
    """Render DDRIntermediate to a Word document (.docx)."""
    from docx import Document
    from docx.shared import Inches

    out_path = Path(out_path).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    _base = Path(base_dir).resolve() if base_dir else Path.cwd()

    doc = Document()
    doc.add_heading("DETAILED DIAGNOSTIC REPORT (DDR)", 0)

    doc.add_heading("1. Property Issue Summary", level=1)
    doc.add_paragraph((ddr.property_summary or "").strip() or "Not Available")

    doc.add_heading("2. Area-wise Observations", level=1)
    for area in ddr.areas:
        doc.add_heading(f"Area: {area.area_name or 'Not Available'}", level=2)
        doc.add_paragraph(f"Observation: {(area.observation or '').strip() or 'Not Available'}")
        doc.add_paragraph(f"Thermal Finding: {(area.thermal_finding or '').strip() or 'Not Available'}")
        doc.add_paragraph(f"Probable Root Cause: {(area.root_cause or '').strip() or 'Not Available'}")
        sev = (area.severity or "").strip() or "Not Available"
        reason = (area.severity_reason or "").strip() or "Not Available"
        doc.add_paragraph(f"Severity: {sev} — Reason: {reason}")
        doc.add_paragraph(
            f"Recommended Action: {(area.recommended_action or '').strip() or 'Not Available'}"
        )
        if (area.conflict_note or "").strip():
            doc.add_paragraph(f"Conflict: {area.conflict_note.strip()}")

        doc.add_paragraph("Image:")
        imgs_added = False
        w = Inches(max_image_width_inches)
        for pth in area.image_paths:
            path = Path(pth)
            if not path.is_absolute():
                path = (_base / path).resolve()
            if path.is_file():
                try:
                    doc.add_picture(str(path), width=w)
                    imgs_added = True
                except Exception:
                    continue
        if not imgs_added:
            note = "Image Not Available"
            if area.image_paths:
                note += " (paths present but files not found or unsupported)"
            doc.add_paragraph(note)

    doc.add_page_break()
    doc.add_heading("3. Probable Root Cause Summary", level=1)
    for it in ddr.root_causes_summary or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.root_causes_summary:
        doc.add_paragraph("Not Available")

    doc.add_heading("4. Severity Assessment Summary", level=1)
    doc.add_paragraph("High Severity Issues:")
    for it in ddr.high_severity or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.high_severity:
        doc.add_paragraph("Not Available")
    doc.add_paragraph("Medium Severity Issues:")
    for it in ddr.medium_severity or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.medium_severity:
        doc.add_paragraph("Not Available")
    doc.add_paragraph("Low Severity Issues:")
    for it in ddr.low_severity or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.low_severity:
        doc.add_paragraph("Not Available")

    doc.add_heading("5. Recommended Actions Summary", level=1)
    doc.add_paragraph("Immediate Actions:")
    for it in ddr.immediate_actions or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.immediate_actions:
        doc.add_paragraph("Not Available")
    doc.add_paragraph("Preventive Actions:")
    for it in ddr.preventive_actions or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.preventive_actions:
        doc.add_paragraph("Not Available")
    doc.add_paragraph("Further Investigation Required:")
    for it in ddr.further_investigation or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.further_investigation:
        doc.add_paragraph("Not Available")

    doc.add_heading("6. Additional Notes", level=1)
    for it in ddr.additional_notes or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.additional_notes:
        doc.add_paragraph("Not Available")
    if ddr.conflicts:
        doc.add_paragraph("Document-level conflicts:")
        for it in ddr.conflicts:
            doc.add_paragraph(it, style="List Bullet")

    doc.add_heading("7. Missing or Unclear Information", level=1)
    for it in ddr.missing_information or []:
        doc.add_paragraph(it, style="List Bullet")
    if not ddr.missing_information:
        doc.add_paragraph("Not Available")

    doc.save(str(out_path))
